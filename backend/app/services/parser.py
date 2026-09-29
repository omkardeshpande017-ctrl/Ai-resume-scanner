import re
from pathlib import Path
import fitz
from docx import Document

ALLOWED = {".pdf", ".docx"}

def extract_text(path: str) -> str:
    ext = Path(path).suffix.lower()
    if ext == ".pdf":
        doc = fitz.open(path)
        try:
            text = "\n".join(page.get_text("text") for page in doc)
        finally:
            doc.close()
        return text
    if ext == ".docx":
        doc = Document(path)
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    raise ValueError("Unsupported file format")

def clean_text(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()

def _first(pattern, text):
    m = re.search(pattern, text, re.I | re.M)
    return m.group(0).strip() if m else ""

def parse_contact(text: str):
    emails = re.findall(r"[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}", text, re.I)
    phones = re.findall(r"(?:\+?\d[\d\s().-]{8,}\d)", text)
    urls = re.findall(r"https?://[^\s)]+", text, re.I)
    lines = [x.strip() for x in text.splitlines() if x.strip()]
    name = lines[0] if lines and len(lines[0]) < 80 and not re.search(r"@|resume|curriculum", lines[0], re.I) else ""
    links = {"linkedin":"", "github":"", "portfolio":""}
    for u in urls:
        low=u.lower()
        if "linkedin.com" in low: links["linkedin"] = u
        elif "github.com" in low: links["github"] = u
        else: links["portfolio"] = u
    return {"name": name, "email": emails[0] if emails else "", "phone": phones[0] if phones else "", **links}

def sections(text: str):
    headings = ["summary","objective","skills","technical skills","education","experience","work experience","internship","internships","projects","certifications","achievements","awards"]
    lines = text.splitlines()
    out = {h: [] for h in headings}
    current = None
    for line in lines:
        normalized = re.sub(r"[^a-z ]", "", line.lower()).strip()
        match = next((h for h in headings if normalized == h), None)
        if match:
            current = match; continue
        if current and line.strip(): out[current].append(line.strip())
    return {k: "\n".join(v) for k,v in out.items() if v}

def parse_resume(text: str):
    contact = parse_contact(text)
    secs = sections(text)
    education = secs.get("education", "")
    skills_blob = " ".join(secs.get(k, "") for k in ["skills","technical skills"])
    return {**contact, "location":"", "education":education, "skills_raw":skills_blob, "sections":secs, "raw_text":text}
